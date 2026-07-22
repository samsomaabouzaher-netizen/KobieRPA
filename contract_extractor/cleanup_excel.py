import openpyxl
import re
import os
from main import load_or_create_config

def clean_excel():
    print("Starte Bereinigung...")
    config = load_or_create_config()
    excel_path = config["excel_pfad"]
    
    if not os.path.exists(excel_path):
        print(f"Datei nicht gefunden: {excel_path}")
        return

    wb = openpyxl.load_workbook(excel_path, data_only=False)
    ws = wb.active
    
    headers = [cell.value for cell in ws[1]]
    if "Kreditinstitut" not in headers:
        print("Spalte Kreditinstitut nicht gefunden!")
        return
        
    kredit_idx = headers.index("Kreditinstitut") + 1
    zahlung_idx = headers.index("Zahlungspflichtiger") + 1 if "Zahlungspflichtiger" in headers else None
    
    changes = 0
    for row in range(2, ws.max_row + 1):
        kredit_val = ws.cell(row=row, column=kredit_idx).value
        
        if kredit_val and isinstance(kredit_val, str):
            kredit_val_clean = kredit_val.strip()
            
            # Prüfen ob es eine IBAN ist (fängt mit DE an und hat Zahlen, auch mit Leerzeichen)
            iban_pattern = re.compile(r'^DE\s*\d{2}\s*')
            if iban_pattern.match(kredit_val_clean):
                print(f"Zeile {row}: Falsche IBAN im Kreditinstitut gefunden und gelöscht -> {kredit_val}")
                ws.cell(row=row, column=kredit_idx).value = ""
                changes += 1
                continue
                
            # Prüfen ob versehentlich der Name des Zahlungspflichtigen eingetragen wurde
            if zahlung_idx:
                zahlung_val = ws.cell(row=row, column=zahlung_idx).value
                if zahlung_val and isinstance(zahlung_val, str) and kredit_val_clean.lower() == zahlung_val.strip().lower():
                    print(f"Zeile {row}: Name im Kreditinstitut gefunden und gelöscht -> {kredit_val}")
                    ws.cell(row=row, column=kredit_idx).value = ""
                    changes += 1
                    continue
                    
    wb.save(excel_path)
    print(f"Bereinigung abgeschlossen! {changes} falsche Einträge im Kreditinstitut wurden entfernt.")

if __name__ == "__main__":
    clean_excel()
